from __future__ import annotations

import csv
import io
from datetime import datetime
from typing import Any

from openpyxl import load_workbook

from app.services.validation_service import ValidationService


class UploadService:
    @staticmethod
    def _coerce_numeric(value: Any) -> Any:
        if value is None or value == "":
            return None
        if isinstance(value, bool):
            return float(value)
        if isinstance(value, (int, float)):
            return float(value)
        if isinstance(value, str):
            cleaned = value.strip().replace(",", "")
            if cleaned.lower() in {"", "nan", "null", "none"}:
                return None
            try:
                return float(cleaned)
            except ValueError:
                return value
        return value

    @staticmethod
    def _normalize_row(row: dict[str, Any]) -> dict[str, Any]:
        normalized = {}
        for key, value in row.items():
            if value is None:
                continue
            normalized[str(key).strip()] = value

        aliases = {
            "date": "timestamp",
            "time": "timestamp",
            "datetime": "timestamp",
            "ts": "timestamp",
            "generation": "generation_mw",
            "power": "power_mw",
            "energy": "energy_mwh",
            "soc": "battery_soc_percent",
            "soC": "battery_soc_percent",
            "state_of_charge": "battery_soc_percent",
            "status": "asset_status",
        }
        for old_key, new_key in aliases.items():
            if old_key in normalized and new_key not in normalized:
                normalized[new_key] = normalized.pop(old_key)

        if "generation_mw" not in normalized and "power_mw" in normalized:
            normalized["generation_mw"] = normalized["power_mw"]
        if "timestamp" in normalized and not isinstance(normalized["timestamp"], str):
            normalized["timestamp"] = normalized["timestamp"].isoformat()

        numeric_keys = {
            "generation_mw",
            "power_mw",
            "energy_mwh",
            "battery_soc_percent",
            "wind_speed_m_s",
            "wind_direction_deg",
            "temperature_c",
            "cloud_cover_percent",
            "irradiance_w_m2",
            "charge_power_mw",
            "discharge_power_mw",
            "capacity_mw",
        }
        for key in numeric_keys:
            if key in normalized:
                normalized[key] = UploadService._coerce_numeric(normalized[key])

        return normalized

    @staticmethod
    def parse_csv(file_bytes: bytes) -> list[dict[str, Any]]:
        text_buffer = io.StringIO(file_bytes.decode("utf-8-sig"))
        reader = csv.DictReader(text_buffer)
        return [UploadService._normalize_row(row) for row in reader]

    @staticmethod
    def parse_excel(file_bytes: bytes) -> list[dict[str, Any]]:
        workbook = load_workbook(filename=io.BytesIO(file_bytes), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        data_rows = []
        for row in rows[1:]:
            data = {}
            for idx, value in enumerate(row):
                if idx < len(headers):
                    key = headers[idx] or f"column_{idx + 1}"
                    data[key] = value
            data_rows.append(UploadService._normalize_row(data))
        return data_rows

    @staticmethod
    def process_upload(file_bytes: bytes, *, filename: str, asset_type: str = "solar", capacity_mw: float = 10.0) -> dict[str, Any]:
        lower_name = filename.lower()
        if lower_name.endswith(".csv"):
            rows = UploadService.parse_csv(file_bytes)
        elif lower_name.endswith((".xlsx", ".xls")):
            rows = UploadService.parse_excel(file_bytes)
        else:
            raise ValueError("Unsupported file type; upload a CSV or Excel file")

        validation = ValidationService.validate_rows(
            [{**row, "asset_type": asset_type} for row in rows],
            capacity_mw=capacity_mw,
        )

        return {
            "source": "file-upload",
            "filename": filename,
            "accepted_rows": validation["accepted_rows"],
            "rejected_rows": validation["rejected_rows"],
            "validation_errors": validation["validation_errors"],
            "duplicate_count": validation["duplicate_count"],
            "processing_summary": validation["processing_summary"],
            "asset_type": asset_type,
            "capacity_mw": capacity_mw,
        }
